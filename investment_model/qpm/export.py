"""CSV and Excel exports."""
from __future__ import annotations
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", start_color="22304A")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BASE = Font(name="Arial", size=10)
INPUT = Font(name="Arial", size=10, color="0000FF")
PCT_HINT = ("target", "actual", "weight", "income", "yield", "ret", "vol", "mdd", "cvar", "cc", "credit", "clo", "equity",
            "defensive", "current", "raw", "active", "dev", "band", "loss", "turnover", "delta", "share", "carry", "limit", "final", "post")


def scores_df(result):
    return pd.DataFrame(result["scores"]).set_index("ticker")


def target_df(result):
    return pd.DataFrame(result["portfolio"]["rows"]).set_index("ticker")


def macro_df(result):
    return pd.DataFrame(result["macro"]["factors"]).set_index("factor")


def _write(ws, df, start_row=1, pct_cols=None, title=None):
    r0 = start_row
    if title:
        ws.cell(row=r0, column=1, value=title).font = Font(name="Arial", size=11, bold=True)
        r0 += 1
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=r0, column=j, value=str(c))
        cell.fill, cell.font, cell.alignment = HDR_FILL, HDR_FONT, Alignment(horizontal="center", wrap_text=True)
    pc = set(pct_cols) if pct_cols is not None else {c for c in cols if any(h in str(c).lower() for h in PCT_HINT)}
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (list, dict)):
                v = str(v)
            try:
                if pd.isna(v):
                    v = None
            except (TypeError, ValueError):
                pass
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BASE
            if isinstance(v, float):
                cell.number_format = "0.00%" if cols[j - 1] in pc else "0.000"
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(str(c)) + 4))
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return r0 + len(df) + 2


def export_csv_bundle(result, outdir):
    outdir.mkdir(parents=True, exist_ok=True)
    m = result["meta"]["month"]
    files = {}
    files["target"] = outdir / f"target_portfolio_{m}.csv"
    target_df(result)[["name", "asset_class", "target", "model_current", "raw_opt", "yield_ttm", "score", "price"]].to_csv(files["target"], encoding="utf-8-sig")
    files["scores"] = outdir / f"etf_scores_{m}.csv"
    scores_df(result).to_csv(files["scores"], encoding="utf-8-sig")
    files["macro"] = outdir / f"macro_{m}.csv"
    macro_df(result).to_csv(files["macro"], encoding="utf-8-sig")
    if result.get("rebalance"):
        files["rebalance"] = outdir / f"rebalancing_order_{m}.csv"
        pd.DataFrame(result["rebalance"]["rows"]).to_csv(files["rebalance"], index=False, encoding="utf-8-sig")
        files["actual"] = outdir / f"actual_portfolio_{m}.csv"
        pd.DataFrame(result["actual"]["rows"]).to_csv(files["actual"], index=False, encoding="utf-8-sig")
    return files


def export_excel(result, path, index=None):
    wb = Workbook()
    meta, m = result["meta"], result["metrics"]
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = f"TAA Income Quant PM — {meta['month']} 데이터 → {meta['target_month']} 목표"
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    info = [("Data As Of", meta["data_as_of"]), ("Rebalance Date", meta["rebalance_date"]), ("Model / System", f"{meta['model_version']} / {meta['system_version']}"),
            ("Macro Regime", result["macro"]["composite_long"]), ("Risk / Defensive State", f"{result['macro']['risk_state']} / {result['macro']['def_state']}"),
            ("Current basis", meta["current_basis"]), ("Validation", result["validation"]["summary"]["status"])]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=i, column=2, value=v).font = BASE
    keys = ["IncomeYield", "Income_3M", "Income_6M", "Income_12M", "ExpTotalReturn", "Vol_36M", "Vol_60M", "Sharpe_36M", "Sortino_36M", "MDD_60M",
            "MDD_Full", "CVaR95m_60M", "CVaR95m_Full", "EffEquityBeta", "Duration", "CoveredCall", "CreditRisk", "CLO_AAA", "PureEquity", "Defensive", "Turnover_vs_model_current"]
    mdf = pd.DataFrame({"Metric": keys, "Value": [m.get(k) for k in keys]})
    _write(ws, mdf, start_row=12, pct_cols={"Value"})
    for r in range(13, 13 + len(keys)):
        if ws.cell(row=r, column=1).value in ("Sharpe_36M", "Sortino_36M", "EffEquityBeta", "Duration"):
            ws.cell(row=r, column=2).number_format = "0.00"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 44
    _write(wb.create_sheet("Target Portfolio"), target_df(result).reset_index()[["ticker", "name", "asset_class", "target", "model_current", "raw_opt", "unconstrained", "yield_ttm", "exp_ret", "score", "price", "beta_eq", "duration"]])
    wsA = wb.create_sheet("Actual Portfolio")
    wsR = wb.create_sheet("Rebalancing")
    if result.get("rebalance"):
        _write(wsA, pd.DataFrame(result["actual"]["rows"]))
        rb = result["rebalance"]
        wsR["A1"], wsR["A2"], wsR["A3"] = "Portfolio Value", "Rebalance Band", "Transaction cost (bps)"
        wsR["B1"], wsR["B2"], wsR["B3"] = rb["totals"]["portfolio_value"], rb["band"], rb["tc_bps"]
        for c in ("B1", "B2", "B3"):
            wsR[c].font = INPUT
        wsR["B1"].number_format, wsR["B2"].number_format = "#,##0", "0.0%"
        hdr = ["ETF", "Class", "Target", "Actual", "Difference", "Action", "Trade Amount", "Price", "Shares"]
        for j, h in enumerate(hdr, 1):
            c = wsR.cell(row=5, column=j, value=h)
            c.fill, c.font = HDR_FILL, HDR_FONT
        uni = [r for r in rb["rows"] if r["asset_class"] != "유니버스 외"]
        for i, r in enumerate(uni, 6):
            vals = [r["ticker"], r["asset_class"], r["target"], r["actual"]]
            for j, v in enumerate(vals, 1):
                wsR.cell(row=i, column=j, value=v).font = BASE
            wsR.cell(row=i, column=3).font = INPUT
            wsR.cell(row=i, column=4).font = INPUT
            wsR.cell(row=i, column=5, value=f"=D{i}-C{i}")
            wsR.cell(row=i, column=6, value=f'=IF(E{i}>$B$2,"SELL",IF(E{i}<-$B$2,"BUY","HOLD"))')
            wsR.cell(row=i, column=7, value=f'=IF(F{i}="HOLD",0,(C{i}-D{i})*$B$1)')
            wsR.cell(row=i, column=8, value=r["price"]).font = INPUT
            wsR.cell(row=i, column=9, value=f"=IF(H{i}>0,TRUNC(G{i}/H{i}),0)")
            for j in (3, 4, 5):
                wsR.cell(row=i, column=j).number_format = "0.00%"
            wsR.cell(row=i, column=7).number_format = "#,##0;(#,##0);-"
            wsR.cell(row=i, column=8).number_format = "0.00"
        last = 5 + len(uni)
        tot = last + 2
        labels = [("Total BUY", f'=SUMIF(G6:G{last},">0")'), ("Total SELL (universe)", f'=-SUMIF(G6:G{last},"<0")'),
                  ("Est. transaction cost", f"=(B{tot}+B{tot+1})*$B$3/10000"), ("Turnover (universe)", f"=(B{tot}+B{tot+1})/$B$1")]
        for k, (lab, fml) in enumerate(labels):
            wsR.cell(row=tot + k, column=1, value=lab).font = Font(name="Arial", size=10, bold=True)
            wsR.cell(row=tot + k, column=2, value=fml).number_format = "0.00%" if "Turnover" in lab else "#,##0"
        oth = [r for r in rb["rows"] if r["asset_class"] == "유니버스 외"]
        if oth:
            _write(wsR, pd.DataFrame(oth)[["ticker", "actual", "trade_amount"]], start_row=tot + 6, title="유니버스 외 보유 (전량 매도 대상)")
        wsR.column_dimensions["A"].width = 24
        for col in "BCDEFGHI":
            wsR.column_dimensions[col].width = 14
    else:
        wsA["A1"] = "Actual Portfolio: Not Provided"
        wsR["A1"] = "Actual Portfolio: Not Provided — 실보유를 입력하면 주문이 생성됩니다."
    _write(wb.create_sheet("ETF Score"), scores_df(result).reset_index().drop(columns=["hist_z"], errors="ignore"))
    wsM = wb.create_sheet("Macro")
    nxt = _write(wsM, macro_df(result).reset_index()[["factor", "level_fmt", "chg_1m", "chg_3m", "chg_6m", "chg_unit", "z", "pct", "regime", "detail"]], pct_cols=set())
    _write(wsM, pd.DataFrame(result["macro"]["history"]), start_row=nxt + 1, pct_cols=set(), title="최근 36개월 매크로 이력")
    wsK = wb.create_sheet("Risk")
    nxt = _write(wsK, pd.DataFrame(result["constraints"]["rows"])[["name", "op", "limit", "raw", "final", "status", "dual", "note"]], pct_cols={"limit", "raw", "final"}, title="Constraint diagnostics")
    _write(wsK, pd.DataFrame(result["frontier"]["rows"]), start_row=nxt + 1, title="Income-Risk frontier")
    ws_st = wb.create_sheet("Stress Test")
    _write(ws_st, pd.DataFrame([dict(scenario=r["scenario"], desc=r["desc"], loss=r["loss"], factor=str(r["factor"]), top=str(r["top"]), worst=str(r["worst"])) for r in result["stress"]["rows"]]))
    bt = result.get("backtest") or {}
    ws_b = wb.create_sheet("Backtest")
    if bt.get("rows"):
        nxt = _write(ws_b, pd.DataFrame(bt["rows"]), title="Backtest 2017-01~2026-08 (v3.0)")
        if bt.get("calendar"):
            _write(ws_b, pd.DataFrame(bt["calendar"]), start_row=nxt + 1, title="연도별 수익률")
    wf = result.get("walk_forward") or {}
    ws_w = wb.create_sheet("Walk Forward")
    if wf.get("choices"):
        _write(ws_w, pd.DataFrame(wf["choices"]), pct_cols=set(), title="연도별 선택 파라미터")
    ws_s = wb.create_sheet("Parameter Stability")
    _write(ws_s, pd.DataFrame(result["stability"]["rows"]), title=f"Verdict: {result['stability']['verdict']}")
    ws_h = wb.create_sheet("History")
    if index and index.get("months"):
        _write(ws_h, pd.DataFrame([dict(month=x["month"], as_of=x["as_of"], target_month=x["target_month"], regime=x["regime"], income=x["income"], source=x.get("source")) for x in index["months"]]))
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path
